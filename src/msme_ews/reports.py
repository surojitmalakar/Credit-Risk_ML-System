"""Offline PDF and Excel export for a selected financial assessment."""

from __future__ import annotations

from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing


def _display(value: object) -> str:
    if value is None or pd.isna(value):
        return "Not available"
    if isinstance(value, (int, float)):
        return f"{value:,.4g}"
    return str(value)


def _assessment_table(assessment: dict[str, Any]) -> list[list[str]]:
    probability = assessment.get("distress_probability")
    return [
        ["Assessment method", str(assessment.get("method", "Existing ML model"))],
        ["Financial health score", _display(assessment.get("health_score"))],
        ["Risk index / model estimate", f"{probability:.1%}" if probability is not None else "Not available"],
        ["Risk category", str(assessment.get("risk_category", "Not available"))],
        ["Data coverage", str(assessment.get("coverage_label", "Not available"))],
    ]


def create_credit_assessment_pdf(
    company_name: str,
    statement: pd.DataFrame,
    analysis: dict[str, Any],
    assessment: dict[str, Any],
    warnings: list[str],
    recommendations: list[str],
) -> bytes:
    """Create a compact local PDF report using ReportLab."""
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
        title=f"Credit Assessment - {company_name}",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="RiskBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        alignment=TA_LEFT,
        textColor=colors.HexColor("#243449"),
    ))
    story = [
        Paragraph("CREDIT RISK AI", styles["Title"]),
        Paragraph("MSME Financial Assessment", styles["Heading2"]),
        Paragraph(f"Company: {escape(company_name)}", styles["RiskBody"]),
        Spacer(1, 10),
        Paragraph("Assessment summary", styles["Heading3"]),
    ]
    summary = Table(_assessment_table(assessment), colWidths=[2.3 * inch, 4.6 * inch])
    summary.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF1F8")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#152238")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([summary, Spacer(1, 10), Paragraph("Financial ratios", styles["Heading3"])])
    ratio_rows = [["Ratio", "Observed value"]] + [
        [str(label), _display(value)]
        for label, value in analysis["ratios"].items()
        if value is not None
    ]
    ratios = Table(ratio_rows, colWidths=[3.4 * inch, 3.5 * inch], repeatRows=1)
    ratios.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1B2A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("PADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(ratios)
    for title, entries in (
        ("Early warning indicators", warnings or ["No configured warning indicators were triggered."]),
        ("Recommendations", recommendations),
    ):
        story.extend([Spacer(1, 10), Paragraph(title, styles["Heading3"])])
        story.extend(Paragraph(f"- {escape(str(entry))}", styles["RiskBody"]) for entry in entries)

    available = statement.iloc[0].get("period") if not statement.empty else None
    story.extend([
        Spacer(1, 12),
        Paragraph(f"Selected financial period: {escape(_display(available))}", styles["RiskBody"]),
        Spacer(1, 4),
        Paragraph(
            "This report is an analytical research summary, not a lending decision or a guaranteed prediction. "
            "Rule-based results are heuristic indices and are not calibrated probabilities of default.",
            styles["RiskBody"],
        ),
    ])
    document.build(story)
    return output.getvalue()


def create_excel_analysis(
    statement: pd.DataFrame,
    analysis: dict[str, Any],
    assessment: dict[str, Any],
    warnings: list[str],
    recommendations: list[str],
) -> bytes:
    """Create a multi-sheet XLSX analysis workbook using openpyxl."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        statement.to_excel(writer, sheet_name="Financial Data", index=False)
        pd.DataFrame(
            [{"Ratio": label, "Value": value} for label, value in analysis["ratios"].items()]
        ).to_excel(writer, sheet_name="Financial Ratios", index=False)
        pd.DataFrame(
            [{"Metric": key.replace("_", " ").title(), "Value": value} for key, value in assessment.items()
             if key in {"method", "distress_probability", "risk_category", "health_score", "confidence_note", "coverage_label"}]
        ).to_excel(writer, sheet_name="Risk Assessment", index=False)
        pd.DataFrame({"Early Warning Indicator": warnings}).to_excel(
            writer, sheet_name="Early Warnings", index=False,
        )
        pd.DataFrame({"Recommendation": recommendations}).to_excel(
            writer, sheet_name="Recommendations", index=False,
        )
        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="0D1B2A")
        header_font = Font(bold=True, color="FFFFFF")
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
            for column_cells in worksheet.columns:
                width = min(48, max(
                    12,
                    max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0) + 2,
                ))
                worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
                for cell in column_cells[1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return output.getvalue()


def create_data_intelligence_pdf(analysis: dict[str, Any], filename: str) -> bytes:
    """Create an offline exploratory report based only on analyzed upload contents."""
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=letter, rightMargin=0.55 * inch, leftMargin=0.55 * inch,
        topMargin=0.55 * inch, bottomMargin=0.55 * inch,
        title=f"Data Intelligence Report - {filename}",
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("DATA INTELLIGENCE REPORT", styles["Title"]),
        Paragraph(escape(filename), styles["Heading2"]),
        Paragraph(escape(analysis["executive_summary"]), styles["BodyText"]),
        Spacer(1, 10),
        Paragraph("Dataset profile", styles["Heading3"]),
        Paragraph(
            f"Type: {escape(analysis['dataset_type'])}; rows: {len(analysis['data']):,}; "
            f"variables: {analysis['data'].shape[1] - 2:,}; "
            f"data quality: {analysis['data_quality_percent']:.1f}%; "
            f"duplicates: {analysis['duplicate_count']:,}; "
            f"anomalies: {analysis['anomaly_count']:,}. "
            "The Excel processed-data sheet includes up to the first 50,000 rows.",
            styles["BodyText"],
        ),
    ]
    detected = analysis["profile"].loc[
        analysis["profile"]["Detected meaning"].astype(str).ne(""),
        ["Variable", "Detected meaning"],
    ]
    if not detected.empty:
        story.extend([
            Paragraph("Detected variables", styles["Heading3"]),
            Paragraph(escape("; ".join(
                f"{row['Variable']}: {row['Detected meaning']}"
                for _, row in detected.head(20).iterrows()
            )), styles["BodyText"]),
        ])
    model_details = analysis["model"]
    if model_details.get("status") == "evaluated":
        score_lines = [
            f"{key.replace('_', ' ').title()}: {value:.3f}"
            for key, value in model_details.items()
            if key in {"accuracy", "balanced_accuracy", "mae", "r2"} and isinstance(value, (int, float))
        ]
        story.extend([
            Paragraph("Exploratory model results", styles["Heading3"]),
            Paragraph(
                escape(
                    f"{model_details.get('model')} predicting {model_details.get('target')}; "
                    f"{', '.join(score_lines)}. " + str(model_details.get("limitation", ""))
                ),
                styles["BodyText"],
            ),
        ])
        if model_details.get("top_features"):
            story.append(Paragraph(
                "Top model features: " + escape(", ".join(
                    f"{item['feature']} ({item['importance']:.3f})"
                    for item in model_details["top_features"][:8]
                )),
                styles["BodyText"],
            ))
    numeric_profile = analysis["profile"].loc[
        analysis["profile"]["Type"].eq("Numeric"), ["Variable", "Missing %"]
    ].head(8)
    if not numeric_profile.empty:
        chart = VerticalBarChart()
        chart.x = 30
        chart.y = 25
        chart.height = 130
        chart.width = 480
        chart.data = [numeric_profile["Missing %"].astype(float).tolist()]
        chart.categoryAxis.categoryNames = numeric_profile["Variable"].astype(str).tolist()
        chart.valueAxis.valueMin = 0
        chart.valueAxis.valueMax = max(10, float(numeric_profile["Missing %"].max()))
        chart.valueAxis.valueStep = max(1, chart.valueAxis.valueMax / 5)
        drawing = Drawing(500, 180)
        drawing.add(chart)
        story.extend([Spacer(1, 8), Paragraph("Missing values by numeric variable (%)", styles["Heading3"]), drawing])
    for title, values in (
        ("Key patterns", analysis["findings"]),
        ("Early warnings", analysis["early_warnings"] or ["No data-derived early warning was triggered."]),
        ("Risk findings and limitations", [
            "Anomaly flags describe statistical unusualness; they do not establish default, fraud, or misconduct.",
            analysis["model"].get("limitation", analysis["model"].get("reason", "No supervised model was evaluated.")),
        ]),
        ("Recommendations", analysis["recommendations"]),
    ):
        story.append(Paragraph(title, styles["Heading3"]))
        story.extend(Paragraph(f"- {escape(str(value))}", styles["BodyText"]) for value in values)
    if not analysis["correlations"].empty:
        story.extend([Paragraph("Strongest numeric associations", styles["Heading3"])])
        rows = [["Variable A", "Variable B", "Absolute correlation"]]
        rows.extend([
            [str(row["Variable A"]), str(row["Variable B"]), f"{row['Absolute correlation']:.3f}"]
            for _, row in analysis["correlations"].head(10).iterrows()
        ])
        table = Table(rows, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1B2A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
    if not analysis["categorical_summary"].empty:
        story.extend([
            Paragraph("Categorical distributions", styles["Heading3"]),
            Table(
                [["Variable", "Value", "Records", "Share %"]]
                + [
                    [str(row["Variable"]), str(row["Value"]), str(row["Records"]), str(row["Share %"])]
                    for _, row in analysis["categorical_summary"].head(20).iterrows()
                ],
                repeatRows=1,
            ),
        ])
    if not analysis["high_risk_groups"].empty:
        story.extend([Paragraph("Observed high-risk segments", styles["Heading3"])])
        rows = [["Variable", "Group", "Records", "Observed risk pattern %"]]
        rows.extend([
            [
                str(row["Variable"]), str(row["Group"]), str(row["Records"]),
                f"{row['Observed risk pattern %']:.1f}",
            ]
            for _, row in analysis["high_risk_groups"].head(12).iterrows()
        ])
        story.append(Table(rows, repeatRows=1))
    document.build(story)
    return output.getvalue()


def create_data_intelligence_excel(analysis: dict[str, Any]) -> bytes:
    """Create a multi-sheet workbook for profiling, statistics, and risk results."""
    output = BytesIO()
    model = analysis["model"]
    model_rows = [
        {"Metric": key, "Value": value if not isinstance(value, (dict, list)) else str(value)}
        for key, value in model.items()
    ]
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([
            {"Metric": "Dataset type", "Value": analysis["dataset_type"]},
            {"Metric": "Rows", "Value": len(analysis["data"])},
            {"Metric": "Variables", "Value": len(analysis["profile"])},
            {"Metric": "Data quality %", "Value": analysis["data_quality_percent"]},
            {"Metric": "Missing cells %", "Value": analysis["missing_percent"]},
            {"Metric": "Duplicate records", "Value": analysis["duplicate_count"]},
            {"Metric": "Anomalies", "Value": analysis["anomaly_count"]},
            {"Metric": "Processed rows included in workbook", "Value": min(len(analysis["data"]), 50_000)},
        ]).to_excel(writer, sheet_name="Overview", index=False)
        analysis["data"].head(50_000).to_excel(writer, sheet_name="Processed Data", index=False)
        analysis["profile"].to_excel(writer, sheet_name="Data Profile", index=False)
        analysis["statistics"].to_excel(writer, sheet_name="Statistics", index=False)
        analysis["categorical_summary"].to_excel(writer, sheet_name="Categorical Summary", index=False)
        analysis["correlations"].to_excel(writer, sheet_name="Correlations", index=False)
        analysis["trends"].to_excel(writer, sheet_name="Trends", index=False)
        analysis["risk_results"].to_excel(writer, sheet_name="Risk Results", index=False)
        analysis["segments"].to_excel(writer, sheet_name="Segments", index=False)
        analysis["high_risk_groups"].to_excel(writer, sheet_name="Risk Segments", index=False)
        analysis["target_distribution"].to_excel(writer, sheet_name="Target Distribution", index=False)
        pd.DataFrame(model_rows).to_excel(writer, sheet_name="Model Results", index=False)
        pd.DataFrame({"Finding": analysis["findings"]}).to_excel(writer, sheet_name="Key Findings", index=False)
        pd.DataFrame({"Early Warning": analysis["early_warnings"]}).to_excel(writer, sheet_name="Early Warnings", index=False)
        pd.DataFrame({"Recommendation": analysis["recommendations"]}).to_excel(writer, sheet_name="Recommendations", index=False)
        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="0D1B2A")
        header_font = Font(bold=True, color="FFFFFF")
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
            for column_cells in worksheet.columns:
                width = min(48, max(
                    12,
                    max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0) + 2,
                ))
                worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
                for cell in column_cells[1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return output.getvalue()
