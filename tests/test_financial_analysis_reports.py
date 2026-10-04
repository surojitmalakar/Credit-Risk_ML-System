from __future__ import annotations

import io
import zipfile

import pandas as pd
import pdfplumber
from openpyxl import load_workbook

from msme_ews.financial_analysis import (
    analyze_financials,
    has_sufficient_ml_data,
    recommendations_for,
    rule_based_assessment,
)
from msme_ews.reports import create_credit_assessment_pdf, create_excel_analysis


def test_sparse_data_uses_transparent_rules_and_coverage():
    frame = pd.DataFrame({
        "company_id": ["Sparse Co"],
        "Revenue": [1_000_000],
        "Current_Assets": [80_000],
        "Current_Liabilities": [100_000],
    })

    analysis = analyze_financials(frame, 0)
    result = rule_based_assessment(frame, 0, analysis)
    recommendations = recommendations_for(
        frame.iloc[0],
        analysis,
        ["Negative operating cash flow", "Falling liquidity"],
    )

    assert not has_sufficient_ml_data(frame, 0)
    assert analysis["ratios"]["Current Ratio"] == 0.8
    assert analysis["ratios"]["Working Capital"] == -20_000
    assert result["method"].startswith("Transparent rule-based")
    assert result["risk_category"] in {"Moderate Risk", "High Risk"}
    assert result["coverage_percent"] == analysis["coverage_percent"]
    assert any("cash generation" in item for item in recommendations)


def test_excel_and_pdf_reports_are_generated_locally():
    statement = pd.DataFrame({
        "company_id": ["Acme"],
        "period": [pd.Timestamp("2024-03-31")],
        "Revenue": [1_000_000],
        "Net_Profit": [100_000],
    })
    analysis = analyze_financials(statement, 0)
    result = rule_based_assessment(statement, 0, analysis)
    result["health_score"] = 82
    result["coverage_label"] = "2/13 core fields (15%)"

    pdf = create_credit_assessment_pdf(
        "Acme",
        statement,
        analysis,
        result,
        ["Deteriorating margins"],
        ["Review cost drivers."],
    )
    workbook = create_excel_analysis(
        statement,
        analysis,
        result,
        ["Deteriorating margins"],
        ["Review cost drivers."],
    )

    assert pdf.startswith(b"%PDF")
    with pdfplumber.open(io.BytesIO(pdf)) as document:
        report_text = "\n".join(page.extract_text() or "" for page in document.pages)
    assert "Acme" in report_text
    with zipfile.ZipFile(io.BytesIO(workbook)) as archive:
        assert "xl/workbook.xml" in archive.namelist()
        assert b"Financial Ratios" in archive.read("xl/workbook.xml")
    excel = load_workbook(io.BytesIO(workbook), read_only=True)
    assert excel.sheetnames == [
        "Financial Data",
        "Financial Ratios",
        "Risk Assessment",
        "Early Warnings",
        "Recommendations",
    ]


def test_sparse_assessment_does_not_claim_a_probability_without_values():
    frame = pd.DataFrame({"Revenue": [None], "Total_Assets": [1_000]})
    analysis = analyze_financials(frame, 0)
    result = rule_based_assessment(frame, 0, analysis)

    assert result["distress_probability"] is None
    assert result["risk_category"] == "Insufficient Data"
