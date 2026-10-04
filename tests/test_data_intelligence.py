from __future__ import annotations

import io
import zipfile

import pandas as pd
from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from msme_ews.data_intelligence import analyze_dataset
from msme_ews.documents import extract_financial_document
from msme_ews.reports import create_data_intelligence_excel, create_data_intelligence_pdf


def _customer_master() -> pd.DataFrame:
    rows = []
    for index in range(80):
        flagged = index % 4 == 0
        rows.append({
            "member_ref": f"MEM-{index:04d}",
            "monthly_takehome": 18_000 + (index * 1_250),
            "open_balance": 90_000 + (index * 4_200) + (350_000 if flagged else 0),
            "bureau_points": 520 + (index * 4),
            "joined_on": f"2024-{index % 12 + 1:02d}-15",
            "repayment_outcome": "Late" if flagged else "On time",
            "service_zone": ["North", "South", "West", "East"][index % 4],
        })
    rows[3]["monthly_takehome"] = None
    return pd.DataFrame(rows)


def test_customer_master_is_automatically_profiled_and_evaluated():
    analysis = analyze_dataset(_customer_master(), "customer_master.csv")

    assert analysis["dataset_type"] == "Mixed dataset"
    assert analysis["data"].shape[0] == 80
    assert analysis["missing_percent"] > 0
    assert "member_ref" in analysis["identifiers"]
    assert "repayment_outcome" in analysis["targets"]
    assert analysis["model"]["status"] == "evaluated"
    assert analysis["anomaly_count"] > 0
    assert not analysis["segments"].empty
    assert not analysis["high_risk_groups"].empty
    assert not analysis["trends"].empty
    assert any("association" in finding for finding in analysis["findings"])
    assert analysis["model_evaluation_available"] is True
    assert analysis["analysis_result_title"] == "Supervised Model Evaluation"


def test_unfamiliar_numeric_and_categorical_columns_are_not_rejected():
    source = pd.DataFrame({
        "x1": [1, 3, 5, 7, 9, 11],
        "x2": [2, 4, 8, 16, 32, 64],
        "cohort": ["a", "a", "b", "b", "c", "c"],
    })

    analysis = analyze_dataset(source, "mystery.csv")

    assert analysis["dataset_type"] in {"Unknown dataset", "Transaction dataset"}
    assert analysis["numeric_columns"] == ["x1", "x2"]
    assert len(analysis["correlations"]) == 1
    assert analysis["data_quality_percent"] == 100
    assert "No explicit default target was found." in analysis["executive_summary"]
    assert analysis["model_evaluation_available"] is False
    assert analysis["analysis_result_title"] == "Anomaly Analysis"
    assert analysis["alternative_analysis"] == "Anomaly + Statistical Risk Pattern Detection"
    assert analysis["risk_prediction_summary"] == "Not available without a target."
    assert "No explicit default target was found. Anomaly and risk-pattern analysis was performed instead." in analysis["executive_summary"]


def test_indian_currency_units_and_year_periods_drive_financial_metrics():
    source = pd.DataFrame({
        "Financial Year": [2022, 2023, 2024],
        "Annual Sales": ["₹1,00,000", "Rs. 2 Lakhs", "3 Lakhs"],
        "PAT": ["₹10,000", "Rs. 20,000", "₹30,000"],
        "Total Assets": ["5 Lakhs", "6 Lakhs", "7 Lakhs"],
        "Total Liabilities": ["2 Lakhs", "2 Lakhs", "2 Lakhs"],
        "Cash Balance": ["1 Thousand", "2 Thousand", "3 Crores"],
    })

    analysis = analyze_dataset(source, "financials.csv")
    metrics = analysis["financial_metrics"].set_index("Metric")

    assert analysis["date_columns"] == ["Financial Year"]
    assert analysis["numeric_columns"] == [
        "Annual Sales", "PAT", "Total Assets", "Total Liabilities", "Cash Balance"
    ]
    assert analysis["data"]["Cash Balance"].iloc[-1] == 30_000_000
    assert metrics.loc["Revenue", "Value"] == 200_000
    assert metrics.loc["Revenue growth", "Value"] == 0.75
    assert metrics.loc["Net profit", "Value"] == 20_000
    assert metrics.loc["Profit margin", "Value"] == 0.1


def test_generic_customer_status_is_not_assumed_to_be_a_risk_target():
    source = pd.DataFrame({
        "customer_status": ["Active", "Inactive"] * 20,
        "age": list(range(20, 60)),
        "monthly_spend": list(range(100, 500, 10)),
    })

    analysis = analyze_dataset(source, "customers.csv")

    assert analysis["targets"] == []
    assert analysis["model"]["status"] == "not trained"


def test_numeric_target_is_evaluated_as_regression():
    source = pd.DataFrame({
        "driver": range(60),
        "target": [float(index * 1.75 + (index % 3)) for index in range(60)],
    })

    analysis = analyze_dataset(source, "outcomes.csv")

    assert analysis["targets"] == ["target"]
    assert analysis["model"]["status"] == "evaluated"
    assert analysis["model"]["model"] == "RandomForestRegressor"
    assert "mae" in analysis["model"]


def test_nonfinancial_pdf_text_remains_available_for_general_profiling():
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.drawString(60, 740, "Service quality summary")
    pdf.drawString(60, 710, "Region North active 125")
    pdf.drawString(60, 680, "Region South active 96")
    pdf.save()

    extracted = extract_financial_document(buffer.getvalue(), "operations.pdf")
    analysis = analyze_dataset(extracted.frame, "operations.pdf")

    assert not extracted.frame.empty
    assert analysis["dataset_type"] in {"Unknown dataset", "MSME/business dataset"}
    assert extracted.document_type == "unknown_dataset"


def test_excel_with_unfamiliar_columns_is_loaded_for_general_profiling():
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame({
            "member_ref": ["A-1", "A-2", "A-3"],
            "zone": ["North", "South", "West"],
            "monthly_budget": [1200, 950, 1800],
        }).to_excel(writer, sheet_name="Members", index=False)

    extracted = extract_financial_document(buffer.getvalue(), "membership.xlsx")
    analysis = analyze_dataset(extracted.frame, "membership.xlsx")

    assert extracted.document_type == "unknown_dataset"
    assert analysis["dataset_type"] == "Unknown dataset"
    assert len(extracted.frame) == 3


def test_data_intelligence_reports_include_actual_profile_and_findings():
    analysis = analyze_dataset(_customer_master(), "customer_master.csv")

    pdf = create_data_intelligence_pdf(analysis, "customer_master.csv")
    workbook = create_data_intelligence_excel(analysis)

    assert pdf.startswith(b"%PDF")
    with zipfile.ZipFile(io.BytesIO(workbook)) as archive:
        workbook_xml = archive.read("xl/workbook.xml")
        assert b"Data Profile" in workbook_xml
        assert b"Risk Results" in workbook_xml
        assert b"Processed Data" in workbook_xml
        assert b"Trends" in workbook_xml
    excel = load_workbook(io.BytesIO(workbook), read_only=False)
    assert "Charts" in excel.sheetnames
    assert len(excel["Charts"]._charts) == 1


def test_no_target_excel_report_uses_user_facing_analysis_labels():
    analysis = analyze_dataset(pd.DataFrame({
        "client_ref": ["A-1", "A-2", "A-3"],
        "monthly_budget": [1200, 950, 1800],
        "zone": ["North", "South", "West"],
    }), "clients.csv")
    workbook = create_data_intelligence_excel(analysis)

    with zipfile.ZipFile(io.BytesIO(workbook)) as archive:
        workbook_text = b" ".join(
            archive.read(name) for name in archive.namelist()
            if name.endswith(".xml")
        ).decode("utf-8", errors="replace").lower()

    assert "anomaly analysis" in workbook_text
    assert "not available without a target" in workbook_text
    assert "not trained" not in workbook_text
